from docx import Document
from pypdf import PdfReader


def extract_pdf_text(uploaded_file):
    """Extract text from every page of a PDF file."""
    uploaded_file.seek(0)

    reader = PdfReader(uploaded_file)
    pages_text = []

    for page in reader.pages:
        page_text = page.extract_text() or ""
        pages_text.append(page_text)

    return "\n".join(pages_text)


def extract_docx_text(uploaded_file):
    """Extract text from DOCX paragraphs and tables."""
    uploaded_file.seek(0)

    document = Document(uploaded_file)
    text_parts = []

    # Extract paragraphs.
    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            text_parts.append(paragraph_text)

    # Extract table content.
    for table in document.tables:
        for row in table.rows:
            row_text = " ".join(
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            )

            if row_text:
                text_parts.append(row_text)

    return "\n".join(text_parts)


def extract_resume_text(uploaded_file):
    """Extract text according to the uploaded file extension."""
    if uploaded_file is None:
        raise ValueError("Please upload a resume first.")

    file_name = uploaded_file.name.lower()

    if file_name.endswith(".pdf"):
        return extract_pdf_text(uploaded_file)

    if file_name.endswith(".docx"):
        return extract_docx_text(uploaded_file)

    raise ValueError(
        "Unsupported file format. Upload a PDF or DOCX file."
    )